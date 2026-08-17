"""
Coding 模式 API 路由。
提供文件树、Git 操作、AST 搜索和 Diff 接口。
"""
import html as _html_module
import os
import re
import stat
from pathlib import Path
from typing import Any, BinaryIO, Dict, Iterator, NoReturn, Optional
from urllib.parse import quote as _url_quote

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from loguru import logger
from pydantic import BaseModel
from sqlalchemy.orm import Session
from starlette.background import BackgroundTask

from api.dependencies import get_current_user
from config.settings import settings
from core.coding.file_tree import FileTreeService
from core.coding.git_integration import GitIntegration
from core.coding.ast_search import ASTSearchService
from core.coding.diff_engine import DiffEngine
from db.models import User, get_db
from security.sandbox import is_path_allowed
from workbench.errors import (
    ProjectDisabled,
    ProjectNotFound,
    ProjectRootChanged,
    ProjectRootForbidden,
    ProjectRootInvalid,
    WorkbenchError,
)
from workbench.path_policy import WorkbenchPathPolicy
from workbench.project_service import WorkbenchProjectService

router = APIRouter(prefix="/api/coding", tags=["coding"])

_LEGACY_PROJECT_PATH_SUNSET = "2026-09-01"

# ---- 敏感文件 deny 列表（P0-4）----
# 敏感文件扩展名：命中即拒绝读取/预览，避免泄露密钥与凭证
_DENY_FILE_EXTENSIONS: frozenset[str] = frozenset({
    ".env",
    ".key",        # 私钥
    ".pem",        # 证书/私钥
    ".pfx",        # PKCS#12 证书
    ".p12",        # PKCS#12 证书
    ".keystore",   # Java 密钥库
    ".htpasswd",    # Apache 密码文件
    ".netrc",      # 网络凭证
    ".npmrc",      # npm 认证令牌
    ".pypirc",     # PyPI 认证令牌
    ".gitconfig",  # Git 全局配置（可能含 token）
})

# 敏感文件名（basename 精确匹配）：命中即拒绝
_DENY_FILE_NAMES: frozenset[str] = frozenset({
    ".env",
    ".env.local",
    ".env.production",
    ".env.development",
    ".env.staging",
    "credentials.json",       # 云凭证（GCP/AWS）
    "secrets.json",
    "secrets.yaml",
    "secrets.yml",
    "id_rsa",                 # SSH 私钥
    "id_dsa",
    "id_ecdsa",
    "id_ed25519",
    "authorized_keys",
    "known_hosts",
    ".htpasswd",
    ".netrc",
    ".npmrc",
    ".pypirc",
    "wallet.dat",             # 加密货币钱包
})

# 敏感文件路径正则：匹配即拒绝
_DENY_FILE_PATTERNS: list[re.Pattern[str]] = [
    re.compile(r"(^|[/\\])\.env($|\.[a-zA-Z0-9_-]+$)"),      # .env 及其变体
    re.compile(r"(^|[/\\])\.git[/\\]"),                       # Git 元数据目录
    re.compile(r"(^|[/\\])\.ssh[/\\]"),                        # SSH 目录
    re.compile(r"(^|[/\\])\.aws[/\\]"),                        # AWS 凭证目录
    re.compile(r"(^|[/\\])\.docker[/\\]"),                     # Docker 配置目录
    re.compile(r"(^|[/\\])\.kube[/\\]"),                       # Kubernetes 配置目录
    re.compile(r"_secrets\.(json|ya?ml|ini|toml|env)$", re.IGNORECASE),  # 通用 secrets 文件
]


def _is_sensitive_file(file_path: str) -> bool:
    """判断文件路径是否命中敏感文件 deny 列表。

    检查顺序：扩展名 → 文件名 basename → 路径正则。
    任一命中即返回 True。

    Args:
        file_path: 待检查的文件路径（相对或绝对均可）。

    Returns:
        True 表示该文件为敏感文件，应拒绝访问。
    """
    if not file_path:
        return False

    path_obj = Path(file_path)
    ext = path_obj.suffix.lower()
    if ext in _DENY_FILE_EXTENSIONS:
        return True

    name = path_obj.name.lower()
    if name in _DENY_FILE_NAMES:
        return True

    # 同时检查原始路径与标准化路径，兼容 Windows 反斜杠
    normalized = file_path.replace("\\", "/")
    for pattern in _DENY_FILE_PATTERNS:
        if pattern.search(file_path) or pattern.search(normalized):
            return True

    return False


async def _check_coding_permission(
    current_user: User,
    permission: str,
    db: Session,
) -> None:
    """检查 coding 模块权限。

    管理员自动放行。其他用户一律通过 RBAC 校验：
    - 无显式角色分配时按默认 viewer 角色判定（viewer 无 coding 权限，拒绝）
    - 有显式角色分配但无对应权限时拒绝
    - RBAC 检查异常时 fail-closed：拒绝访问并返回 500

    Args:
        current_user: 当前认证用户。
        permission: 权限标识，如 "coding:read" / "coding:write"。
        db: 数据库会话。

    Raises:
        HTTPException: 权限不足时抛 403；RBAC 检查失败时抛 500。
    """
    # 管理员直接放行
    if getattr(current_user, "role", None) == "admin":
        return

    try:
        from security.rbac import RBACManager

        rbac = RBACManager(db)
        user_id = str(current_user.id)

        # 无显式角色分配时由 check_permission 按默认 viewer 角色判定（无 coding 权限，拒绝）
        allowed = await rbac.check_permission(user_id, permission)
        if not allowed:
            raise HTTPException(
                status_code=403,
                detail=f"权限不足，需要 {permission} 权限或管理员角色"
            )
    except HTTPException:
        raise
    except Exception as exc:
        # RBAC 检查异常时 fail-closed：拒绝访问，不降级放行
        logger.warning(
            f"RBAC 检查异常，已拒绝访问: user_id={current_user.id}, "
            f"permission={permission}, error={exc}"
        )
        raise HTTPException(
            status_code=500,
            detail="RBAC 权限检查失败，已拒绝访问",
        ) from exc

# ---- 可选依赖：Markdown 渲染与净化 ----
# 故意放在模块顶层 try/except 中，避免在 requirements.txt 中强制依赖
try:
    import markdown as _markdown_lib  # type: ignore
except ImportError:
    _markdown_lib = None  # type: ignore

try:
    import bleach as _bleach_lib  # type: ignore
except ImportError:
    _bleach_lib = None  # type: ignore


def get_coding_workbench_path_policy() -> WorkbenchPathPolicy:
    """按当前稳定设置构建 Coding 使用的工作台路径策略。"""
    return WorkbenchPathPolicy.from_settings(settings)


def _coding_error_detail(code: str, message: str) -> dict[str, str]:
    return {"code": code, "message": message}


def _raise_workbench_http_error(exc: WorkbenchError) -> NoReturn:
    """把工作台领域异常映射为 Coding API 的结构化响应。"""
    status_code = 409
    if isinstance(exc, ProjectNotFound):
        status_code = 404
    elif isinstance(exc, ProjectRootForbidden):
        status_code = 403
    elif isinstance(exc, ProjectRootInvalid):
        status_code = 422
    elif isinstance(exc, (ProjectDisabled, ProjectRootChanged)):
        status_code = 409
    raise HTTPException(
        status_code=status_code,
        detail=_coding_error_detail(exc.code, exc.message),
    ) from exc


def _resolve_coding_project_root(
    *,
    project_id: Optional[str],
    project_dir: Optional[str],
    current_user: User,
    db: Session,
    path_policy: WorkbenchPathPolicy,
    legacy_project_path_supplied: bool = False,
) -> str:
    """拒绝旧绝对路径，并将当前用户的项目 ID 解析为本次请求根。"""
    if legacy_project_path_supplied or project_dir is not None:
        raise HTTPException(
            status_code=422,
            detail=_coding_error_detail(
                "legacy_project_path_not_supported",
                "不再支持客户端提供项目路径，请改用 project_id",
            ),
            headers={"Sunset": _LEGACY_PROJECT_PATH_SUNSET},
        )

    normalized_project_id = project_id.strip() if project_id is not None else ""
    if not normalized_project_id:
        raise HTTPException(
            status_code=409,
            detail=_coding_error_detail(
                "workbench_project_required",
                "请先选择工作台项目",
            ),
        )

    service = WorkbenchProjectService(db, path_policy)
    try:
        root = service.resolve_project_root(
            user_id=str(current_user.id),
            user_role=str(current_user.role),
            project_id=normalized_project_id,
        )
    except WorkbenchError as exc:
        _raise_workbench_http_error(exc)
    return str(root)


def get_coding_project_root(
    project_id: Optional[str] = None,
    project_dir: Optional[str] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    path_policy: WorkbenchPathPolicy = Depends(get_coding_workbench_path_policy),
) -> str:
    """为使用查询参数的 Coding 入口提供统一项目根依赖。"""
    return _resolve_coding_project_root(
        project_id=project_id,
        project_dir=project_dir,
        current_user=current_user,
        db=db,
        path_policy=path_policy,
    )


async def get_coding_read_project_root(
    root: str = Depends(get_coding_project_root),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> str:
    """解析项目根后统一验证 Coding 只读权限。"""
    await _check_coding_permission(current_user, "coding:read", db)
    return root


async def get_coding_write_project_root(
    root: str = Depends(get_coding_project_root),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> str:
    """解析项目根后统一验证 Coding 写权限。"""
    await _check_coding_permission(current_user, "coding:write", db)
    return root


def _validate_file_path(file_path: str, project_dir: str, *, is_write: bool = False) -> str:
    """
    校验文件路径在项目目录内，防止路径遍历攻击，并复用沙箱 deny 列表。

    三层防护：
    1. 敏感文件 deny 列表：命中即拒绝（.env、*.key、*.pem 等）
    2. 自有路径遍历防护：确保解析后路径在项目目录内
    3. 沙箱 is_path_allowed：复用全局 deny 规则与 TOCTOU 防护

    Args:
        file_path: 待校验的文件路径（相对项目目录）。
        project_dir: 由工作台统一解析器产生的项目根目录。
        is_write: 是否为写操作，写操作更严格（禁止 glob、禁止 .. 穿越）。

    Returns:
        解析后的绝对路径。

    Raises:
        HTTPException: 路径越权、命中敏感文件或被沙箱拒绝时抛 403。
    """
    # 第 1 层：敏感文件 deny 列表
    if _is_sensitive_file(file_path):
        logger.warning(f"敏感文件访问被拒绝: {file_path}")
        raise HTTPException(status_code=403, detail="禁止访问敏感文件")

    # 第 2 层：自有路径遍历防护（使用 relative_to 替代 startswith，防前缀绕过）
    root_real = Path(os.path.realpath(project_dir)).resolve()
    resolved = Path(os.path.realpath(os.path.join(str(root_real), file_path.lstrip("/\\")))).resolve()
    try:
        resolved.relative_to(root_real)
    except ValueError:
        if resolved != root_real:
            raise HTTPException(status_code=403, detail="禁止访问项目目录外的文件")

    # 第 3 层：复用沙箱 is_path_allowed（含 .env deny 规则、TOCTOU 防护）
    if not is_path_allowed(str(resolved), is_write=is_write, working_dir=str(root_real)):
        logger.warning(f"路径被沙箱拒绝: {file_path} (resolved={resolved})")
        raise HTTPException(status_code=403, detail="路径被沙箱安全策略拒绝")

    return str(resolved)


def _validate_git_file_path(file_path: str, project_dir: str) -> str:
    """拒绝 Git pathspec magic，仅允许项目内普通相对文件路径。"""
    if not file_path or file_path.startswith(":") or Path(file_path).is_absolute():
        raise HTTPException(status_code=403, detail="Git 文件路径无效")
    _validate_file_path(file_path, project_dir, is_write=False)
    return file_path.replace("\\", "/")


def _opened_file_path(file_handle: BinaryIO, fallback_path: str) -> Path:
    """从已打开句柄解析实际目标，避免校验后按路径重新跟随链接。"""
    file_descriptor = file_handle.fileno()
    if os.name == "nt":
        import ctypes
        import msvcrt

        native_handle = msvcrt.get_osfhandle(file_descriptor)
        buffer = ctypes.create_unicode_buffer(32768)
        length = ctypes.windll.kernel32.GetFinalPathNameByHandleW(  # type: ignore[attr-defined]
            native_handle,
            buffer,
            len(buffer),
            0,
        )
        if length == 0 or length >= len(buffer):
            raise OSError("无法解析已打开文件的实际路径")
        normalized = buffer.value
        if normalized.startswith("\\\\?\\"):
            normalized = normalized[4:]
        return Path(normalized).resolve(strict=True)

    proc_path = Path(f"/proc/self/fd/{file_descriptor}")
    if proc_path.exists():
        return Path(os.path.realpath(proc_path)).resolve(strict=True)
    return Path(fallback_path).resolve(strict=True)


def _open_project_binary_file(file_path: str, project_dir: str) -> BinaryIO:
    """打开并复核实际文件句柄仍位于工作台项目根内。"""
    validated_path = _validate_file_path(file_path, project_dir, is_write=False)
    try:
        file_handle = open(validated_path, "rb")
    except OSError as exc:
        raise HTTPException(status_code=404, detail="文件不存在或不是常规文件") from exc
    try:
        opened_path = _opened_file_path(file_handle, validated_path)
        opened_path.relative_to(Path(project_dir).resolve(strict=True))
        if not stat.S_ISREG(os.fstat(file_handle.fileno()).st_mode):
            raise ValueError("打开目标不是常规文件")
    except (OSError, ValueError) as exc:
        file_handle.close()
        raise HTTPException(status_code=403, detail="文件在打开前已离开项目目录") from exc
    return file_handle


def _iter_open_file(file_handle: BinaryIO) -> Iterator[bytes]:
    """流式读取已验证句柄，并在响应结束后关闭。"""
    try:
        while True:
            chunk = file_handle.read(_STREAM_CHUNK_SIZE)
            if not chunk:
                break
            yield chunk
    finally:
        file_handle.close()


# ---- Request Schemas ----

class FileReadRequest(BaseModel):
    path: str = ""
    project_id: Optional[str] = None
    project_dir: Optional[str] = None


class FileWriteRequest(BaseModel):
    path: str
    content: str
    project_id: Optional[str] = None
    project_dir: Optional[str] = None


class SearchRequest(BaseModel):
    pattern: str
    project_id: Optional[str] = None
    project_dir: Optional[str] = None


class GitCommitRequest(BaseModel):
    message: str
    files: Optional[list[str]] = None
    project_id: Optional[str] = None
    project_dir: Optional[str] = None


class DiffRequest(BaseModel):
    original: str
    modified: str


class FileSearchRequest(BaseModel):
    pattern: str
    directory: str = ""
    project_id: Optional[str] = None
    project_dir: Optional[str] = None


# ---- 文件树接口 ----

@router.get("/tree")
def get_file_tree(
    path: str = "",
    root: str = Depends(get_coding_read_project_root),
):
    """
    获取目录树结构。
    """
    _validate_file_path(path, root, is_write=False)
    service = FileTreeService(root)
    return service.get_tree(path)


@router.get("/list")
def list_directory(
    path: str = "",
    root: str = Depends(get_coding_read_project_root),
):
    """
    列出目录内容。
    """
    _validate_file_path(path, root, is_write=False)
    service = FileTreeService(root)
    return service.list_directory(path)


@router.post("/read")
async def read_file(
    body: FileReadRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    path_policy: WorkbenchPathPolicy = Depends(get_coding_workbench_path_policy),
) -> Dict[str, Any]:
    """
    读取文件内容。
    """
    # P0-4: RBAC 权限校验
    await _check_coding_permission(current_user, "coding:read", db)
    root = _resolve_coding_project_root(
        project_id=body.project_id,
        project_dir=body.project_dir,
        current_user=current_user,
        db=db,
        path_policy=path_policy,
        legacy_project_path_supplied="project_dir" in body.model_fields_set,
    )
    # P0-4: 路径校验（含敏感文件 deny 列表与沙箱策略）
    _validate_file_path(body.path, root, is_write=False)
    service = FileTreeService(root)
    result = service.read_file(body.path)
    if "error" in result:
        raise HTTPException(status_code=404, detail=result["error"])
    return result


@router.post("/write")
async def write_file(
    body: FileWriteRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    path_policy: WorkbenchPathPolicy = Depends(get_coding_workbench_path_policy),
) -> Dict[str, Any]:
    """
    写入文件内容。
    """
    # P0-4: RBAC 权限校验
    await _check_coding_permission(current_user, "coding:write", db)
    root = _resolve_coding_project_root(
        project_id=body.project_id,
        project_dir=body.project_dir,
        current_user=current_user,
        db=db,
        path_policy=path_policy,
        legacy_project_path_supplied="project_dir" in body.model_fields_set,
    )
    # P0-4: 路径校验（写操作更严格，禁止 glob 与 .. 穿越）
    _validate_file_path(body.path, root, is_write=True)
    service = FileTreeService(root)
    result = service.write_file(body.path, body.content)
    if "error" in result:
        raise HTTPException(status_code=400, detail=result["error"])
    return result


@router.post("/search-files")
async def search_files(
    body: FileSearchRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    path_policy: WorkbenchPathPolicy = Depends(get_coding_workbench_path_policy),
):
    """
    按文件名搜索。
    """
    root = _resolve_coding_project_root(
        project_id=body.project_id,
        project_dir=body.project_dir,
        current_user=current_user,
        db=db,
        path_policy=path_policy,
        legacy_project_path_supplied="project_dir" in body.model_fields_set,
    )
    await _check_coding_permission(current_user, "coding:read", db)
    _validate_file_path(body.directory, root, is_write=False)
    service = FileTreeService(root)
    results = service.search_files(body.pattern, body.directory)
    return {"results": results, "count": len(results)}


# ---- Git 接口 ----

@router.get("/git/status")
def git_status(root: str = Depends(get_coding_read_project_root)):
    """
    获取 Git 仓库状态。
    """
    git = GitIntegration(root)
    if not git.is_repo():
        return {"error": "不是 Git 仓库", "is_repo": False}
    result = git.get_status()
    result["is_repo"] = True
    return result


@router.get("/git/diff")
def git_diff(
    file_path: Optional[str] = None,
    staged: bool = False,
    root: str = Depends(get_coding_read_project_root),
):
    """
    获取 Git 差异。
    """
    if file_path is not None:
        file_path = _validate_git_file_path(file_path, root)
    git = GitIntegration(root)
    return git.get_diff(file_path, staged)


@router.get("/git/log")
def git_log(
    max_count: int = 20,
    root: str = Depends(get_coding_read_project_root),
):
    """
    获取 Git 提交日志。
    """
    git = GitIntegration(root)
    return git.get_log(max_count)


@router.post("/git/commit")
async def git_commit(
    body: GitCommitRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    path_policy: WorkbenchPathPolicy = Depends(get_coding_workbench_path_policy),
):
    """
    提交 Git 更改。
    """
    root = _resolve_coding_project_root(
        project_id=body.project_id,
        project_dir=body.project_dir,
        current_user=current_user,
        db=db,
        path_policy=path_policy,
        legacy_project_path_supplied="project_dir" in body.model_fields_set,
    )
    await _check_coding_permission(current_user, "coding:write", db)
    files = (
        [_validate_git_file_path(file_path, root) for file_path in body.files]
        if body.files is not None
        else None
    )
    git = GitIntegration(root)
    return git.commit(body.message, files)


@router.get("/git/branches")
def git_branches(root: str = Depends(get_coding_read_project_root)):
    """
    获取分支列表。
    """
    git = GitIntegration(root)
    return git.get_branches()


@router.post("/git/branch")
def git_create_branch(
    name: str,
    root: str = Depends(get_coding_write_project_root),
):
    """
    创建新分支。
    """
    git = GitIntegration(root)
    return git.create_branch(name)


# ---- AST 搜索接口 ----

@router.get("/ast/definitions")
def ast_search_definitions(
    name: str,
    root: str = Depends(get_coding_read_project_root),
):
    """
    搜索函数/类定义。
    """
    service = ASTSearchService(root)
    results = service.search_definitions(name)
    return {"results": results, "count": len(results)}


@router.get("/ast/references")
def ast_search_references(
    name: str,
    root: str = Depends(get_coding_read_project_root),
):
    """
    搜索变量/函数引用。
    """
    service = ASTSearchService(root)
    results = service.search_references(name)
    return {"results": results, "count": len(results)}


@router.post("/ast/search")
async def ast_search_pattern(
    body: SearchRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    path_policy: WorkbenchPathPolicy = Depends(get_coding_workbench_path_policy),
):
    """
    正则模式搜索代码。
    """
    root = _resolve_coding_project_root(
        project_id=body.project_id,
        project_dir=body.project_dir,
        current_user=current_user,
        db=db,
        path_policy=path_policy,
        legacy_project_path_supplied="project_dir" in body.model_fields_set,
    )
    await _check_coding_permission(current_user, "coding:read", db)
    service = ASTSearchService(root)
    results = service.search_pattern(body.pattern)
    return {"results": results, "count": len(results)}


@router.get("/ast/structure")
def ast_get_structure(
    file_path: str,
    root: str = Depends(get_coding_read_project_root),
):
    """
    获取文件结构概览。
    """
    _validate_file_path(file_path, root, is_write=False)
    service = ASTSearchService(root)
    return service.get_structure(file_path)


# ---- Diff 接口 ----

@router.post("/diff")
def compute_diff(body: DiffRequest, current_user=Depends(get_current_user)):
    """
    计算文本差异。
    """
    engine = DiffEngine()
    return engine.compute_inline_diff(body.original, body.modified)


# ---- LSP 接口 ----

class LSPCompletionRequest(BaseModel):
    file_path: str
    line: int
    column: int
    project_id: Optional[str] = None
    project_dir: Optional[str] = None


class LSPHoverRequest(BaseModel):
    file_path: str
    line: int
    column: int
    project_id: Optional[str] = None
    project_dir: Optional[str] = None


@router.get("/lsp/diagnostics")
def get_lsp_diagnostics(
    file_path: str,
    root: str = Depends(get_coding_read_project_root),
):
    """
    获取文件的 LSP 诊断信息（错误/警告）。
    返回语言服务器的可用性和状态信息。
    """
    _validate_file_path(file_path, root)
    try:
        ext = Path(file_path).suffix.lstrip(".").lower()
        lang_map = {"py": "python", "ts": "typescript", "tsx": "typescript", "js": "javascript", "jsx": "javascript"}
        language = lang_map.get(ext, ext)
        from core.coding.lsp_proxy import LSPProxy
        lsp = LSPProxy(root)
        available = lsp.is_available(language)
        return {
            "success": True,
            "language": language,
            "lsp_available": available,
            "diagnostics": [],
            "message": f"LSP 服务器{'可用' if available else '不可用'}: {language}",
        }
    except Exception as exc:
        return {"success": False, "diagnostics": [], "error": str(exc)}


@router.post("/lsp/completions")
async def get_lsp_completions(
    body: LSPCompletionRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    path_policy: WorkbenchPathPolicy = Depends(get_coding_workbench_path_policy),
):
    """
    获取代码补全建议（基于 AST 静态分析）。
    注：完整 LSP 补全需通过前端 Monaco Editor 内置能力实现。
    """
    root = _resolve_coding_project_root(
        project_id=body.project_id,
        project_dir=body.project_dir,
        current_user=current_user,
        db=db,
        path_policy=path_policy,
        legacy_project_path_supplied="project_dir" in body.model_fields_set,
    )
    await _check_coding_permission(current_user, "coding:read", db)
    _validate_file_path(body.file_path, root)
    try:
        from core.coding.ast_search import ASTSearchService
        service = ASTSearchService(root)
        # 基于当前文件提供符号补全
        structure = service.get_structure(body.file_path)
        return {
            "success": True,
            "completions": structure.get("symbols", [])[:20],
            "file_path": body.file_path,
        }
    except Exception as exc:
        return {"success": False, "completions": [], "error": str(exc)}


@router.post("/lsp/hover")
async def get_lsp_hover(
    body: LSPHoverRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    path_policy: WorkbenchPathPolicy = Depends(get_coding_workbench_path_policy),
):
    """
    获取悬停信息（基于 AST 静态分析）。
    注：完整 LSP 悬停需通过前端 Monaco Editor 内置能力实现。
    """
    root = _resolve_coding_project_root(
        project_id=body.project_id,
        project_dir=body.project_dir,
        current_user=current_user,
        db=db,
        path_policy=path_policy,
        legacy_project_path_supplied="project_dir" in body.model_fields_set,
    )
    await _check_coding_permission(current_user, "coding:read", db)
    _validate_file_path(body.file_path, root)
    try:
        from core.coding.ast_search import ASTSearchService
        service = ASTSearchService(root)
        structure = service.get_structure(body.file_path)
        return {
            "success": True,
            "hover": f"文件: {body.file_path}\n行: {body.line}, 列: {body.column}\n符号数: {len(structure.get('symbols', []))}",
            "file_path": body.file_path,
        }
    except Exception as exc:
        return {"success": False, "hover": "", "error": str(exc)}


@router.get("/lsp/symbols")
def get_lsp_symbols(
    file_path: str,
    root: str = Depends(get_coding_read_project_root),
):
    """
    获取文件的符号列表（基于 AST 静态分析）。
    """
    _validate_file_path(file_path, root)
    try:
        from core.coding.ast_search import ASTSearchService
        service = ASTSearchService(root)
        structure = service.get_structure(file_path)
        return {"success": True, "symbols": structure.get("symbols", [])}
    except Exception as exc:
        return {"success": False, "symbols": [], "error": str(exc)}


# ---- Claude Code 模式 ----

class CCModeRequest(BaseModel):
    enabled: bool


# 按用户隔离的 CC 模式状态（user_id -> enabled），防止跨用户状态泄露
_cc_mode_sessions: dict[str, bool] = {}


@router.post("/cc-mode")
def toggle_cc_mode(body: CCModeRequest, current_user: User = Depends(get_current_user)):
    """
    启用或禁用 Claude Code 兼容模式。
    CC 模式下会注入 Coding 专用系统提示和工具定义。
    """
    user_id = str(current_user.id)
    _cc_mode_sessions[user_id] = body.enabled
    return {"success": True, "cc_mode_enabled": body.enabled}


@router.get("/cc-mode")
def get_cc_mode(current_user: User = Depends(get_current_user)):
    """
    获取当前用户的 CC 模式状态。
    """
    user_id = str(current_user.id)
    return {"cc_mode_enabled": _cc_mode_sessions.get(user_id, False)}


# ---- 文件预览接口 ----

# 图片 MIME 映射
_IMAGE_MIME_MAP: dict[str, str] = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
    ".svg": "image/svg+xml",
}

# 音频 MIME 映射
_AUDIO_MIME_MAP: dict[str, str] = {
    ".mp3": "audio/mpeg",
    ".wav": "audio/wav",
    ".ogg": "audio/ogg",
}

# 视频 MIME 映射
_VIDEO_MIME_MAP: dict[str, str] = {
    ".mp4": "video/mp4",
    ".webm": "video/webm",
}

# 文本类 MIME 映射
_TEXT_MIME_MAP: dict[str, str] = {
    ".txt": "text/plain; charset=utf-8",
    ".log": "text/plain; charset=utf-8",
    ".py": "text/x-python; charset=utf-8",
    ".js": "application/javascript; charset=utf-8",
    ".ts": "application/typescript; charset=utf-8",
    ".json": "application/json; charset=utf-8",
    ".html": "text/html; charset=utf-8",
    ".htm": "text/html; charset=utf-8",
    ".css": "text/css; charset=utf-8",
    ".xml": "application/xml; charset=utf-8",
    ".yaml": "application/x-yaml; charset=utf-8",
    ".yml": "application/x-yaml; charset=utf-8",
    ".ini": "text/plain; charset=utf-8",
    ".toml": "text/plain; charset=utf-8",
    ".sh": "application/x-sh; charset=utf-8",
    ".bat": "text/plain; charset=utf-8",
    ".ps1": "text/plain; charset=utf-8",
}

# Office 扩展名集合
_OFFICE_EXTS: set[str] = {".docx", ".xlsx", ".pptx"}

# 流式读取的块大小
_STREAM_CHUNK_SIZE = 64 * 1024  # 64KB


def _download_link_response(
    original_path: str,
    project_id: str,
    error: Optional[str] = None,
) -> dict:
    """生成 Office 文件降级用的下载链接响应体。"""
    payload: dict = {
        "type": "download",
        "url": (
            f"/api/coding/download?path={_url_quote(original_path)}"
            f"&project_id={_url_quote(project_id)}"
        ),
    }
    if error:
        payload["error"] = error
    return payload


def _render_markdown_file(file_handle: BinaryIO) -> dict:
    """从已验证句柄读取 Markdown，并渲染为净化后的 HTML。"""
    try:
        md_text = file_handle.read().decode("utf-8")
    except OSError as exc:
        logger.bind(error_type=type(exc).__name__).opt(exception=True).error(
            "读取 Markdown 文件失败"
        )
        raise HTTPException(status_code=500, detail="读取文件失败，请稍后重试") from exc
    finally:
        file_handle.close()

    if _markdown_lib is None or _bleach_lib is None:
        # 优雅降级：返回 <pre> 包裹的 HTML 转义文本
        escaped = _html_module.escape(md_text)
        return {
            "type": "markdown",
            "html": f"<pre>{escaped}</pre>",
            "error": "markdown or bleach library not installed",
        }

    try:
        raw_html = _markdown_lib.markdown(
            md_text,
            extensions=["extra", "fenced_code", "tables", "sane_lists"],
        )
        cleaned_html = _bleach_lib.clean(
            raw_html,
            tags={
                "h1", "h2", "h3", "h4", "h5", "h6",
                "p", "br", "hr",
                "strong", "em", "b", "i", "u", "s", "del", "ins", "sub", "sup",
                "ul", "ol", "li",
                "blockquote", "code", "pre",
                "a", "img",
                "table", "thead", "tbody", "tfoot", "tr", "th", "td",
                "div", "span", "abbr", "cite", "q",
            },
            attributes={
                "a": ["href", "title", "target", "rel"],
                "img": ["src", "alt", "title", "width", "height"],
                "span": ["class"],
                "div": ["class"],
                "code": ["class"],
                "pre": ["class"],
                "td": ["colspan", "rowspan"],
                "th": ["colspan", "rowspan"],
            },
            protocols=["http", "https", "mailto", "tel"],
            strip=True,
        )
        return {"type": "markdown", "html": cleaned_html}
    except Exception as exc:
        escaped = _html_module.escape(md_text)
        return {
            "type": "markdown",
            "html": f"<pre>{escaped}</pre>",
            "error": f"markdown rendering failed: {exc}",
        }


def _stream_binary_file(
    file_handle: BinaryIO,
    content_type: str,
    request: Request,
    support_range: bool,
) -> StreamingResponse:
    """从同一已验证句柄流式返回文件，并可选支持 Range。"""
    try:
        file_size = os.fstat(file_handle.fileno()).st_size
    except OSError as exc:
        file_handle.close()
        raise HTTPException(status_code=500, detail="读取文件状态失败") from exc

    range_header = request.headers.get("range") if support_range else None
    if range_header:
        try:
            range_spec = range_header.strip()
            if range_spec.startswith("bytes="):
                range_spec = range_spec[len("bytes="):]
            start_str, _, end_str = range_spec.partition("-")
            start = int(start_str) if start_str else 0
            end = int(end_str) if end_str else file_size - 1
            if start < 0 or start >= file_size or end >= file_size or start > end:
                raise ValueError("invalid range")
        except (ValueError, IndexError):
            file_handle.close()
            raise HTTPException(
                status_code=416,
                detail="Range Not Satisfiable",
                headers={"Content-Range": f"bytes */{file_size}"},
            )

        content_length = end - start + 1

        def _iter_range() -> Iterator[bytes]:
            try:
                file_handle.seek(start)
                remaining = content_length
                while remaining > 0:
                    chunk = file_handle.read(min(_STREAM_CHUNK_SIZE, remaining))
                    if not chunk:
                        break
                    yield chunk
                    remaining -= len(chunk)
            finally:
                file_handle.close()

        try:
            return StreamingResponse(
                _iter_range(),
                status_code=206,
                media_type=content_type,
                background=BackgroundTask(file_handle.close),
                headers={
                    "Content-Range": f"bytes {start}-{end}/{file_size}",
                    "Content-Length": str(content_length),
                    "Accept-Ranges": "bytes",
                },
            )
        except Exception:
            file_handle.close()
            raise

    try:
        return StreamingResponse(
            _iter_open_file(file_handle),
            media_type=content_type,
            background=BackgroundTask(file_handle.close),
            headers={
                "Content-Length": str(file_size),
                "Accept-Ranges": "bytes",
            },
        )
    except Exception:
        file_handle.close()
        raise


def _render_office_file(
    file_handle: BinaryIO,
    ext: str,
    original_path: str,
    project_id: str,
) -> dict:
    """渲染 Office 文件。

    优先尝试 mammoth / openpyxl / python-pptx，缺失或失败时降级为下载链接。
    """
    try:
        if ext == ".docx":
            try:
                import mammoth  # type: ignore
            except ImportError:
                return _download_link_response(
                    original_path,
                    project_id,
                    error="mammoth library not installed",
                )
            try:
                file_handle.seek(0)
                result = mammoth.convert_to_html(file_handle)
                return {"type": "markdown", "html": result.value}
            except Exception as exc:
                logger.bind(error_type=type(exc).__name__).opt(exception=True).warning(
                    "Word 文件预览转换失败"
                )
                return _download_link_response(
                    original_path,
                    project_id,
                    error="文档预览失败，请下载后查看",
                )

        if ext == ".xlsx":
            try:
                import openpyxl  # type: ignore
            except ImportError:
                return _download_link_response(
                    original_path,
                    project_id,
                    error="openpyxl library not installed",
                )
            workbook = None
            try:
                file_handle.seek(0)
                workbook = openpyxl.load_workbook(file_handle, read_only=True)
                sheet_names = workbook.sheetnames
                list_html = (
                    "<ul>"
                    + "".join(
                        f"<li>{_html_module.escape(name)}</li>"
                        for name in sheet_names
                    )
                    + "</ul>"
                )
                return {"type": "markdown", "html": list_html}
            except Exception as exc:
                logger.bind(error_type=type(exc).__name__).opt(exception=True).warning(
                    "表格文件预览转换失败"
                )
                return _download_link_response(
                    original_path,
                    project_id,
                    error="表格预览失败，请下载后查看",
                )
            finally:
                if workbook is not None:
                    workbook.close()

        if ext == ".pptx":
            try:
                from pptx import Presentation  # type: ignore
            except ImportError:
                return _download_link_response(
                    original_path,
                    project_id,
                    error="python-pptx library not installed",
                )
            try:
                file_handle.seek(0)
                presentation = Presentation(file_handle)
                slide_count = len(presentation.slides)
                return {
                    "type": "markdown",
                    "html": f"<p>共 {slide_count} 张幻灯片</p>",
                }
            except Exception as exc:
                logger.bind(error_type=type(exc).__name__).opt(exception=True).warning(
                    "演示文件预览转换失败"
                )
                return _download_link_response(
                    original_path,
                    project_id,
                    error="演示文件预览失败，请下载后查看",
                )

        return _download_link_response(original_path, project_id)
    finally:
        file_handle.close()


@router.get("/preview/file")
async def preview_file(
    path: str,
    request: Request,
    project_id: Optional[str] = None,
    project_dir: Optional[str] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    path_policy: WorkbenchPathPolicy = Depends(get_coding_workbench_path_policy),
) -> Dict[str, Any]:
    """
    预览文件内容：按扩展名返回适当格式。

    - Markdown (.md/.markdown)：渲染为 HTML 并经 bleach 净化
    - 图片 (.png/.jpg/.jpeg/.gif/.webp/.svg)：StreamingResponse + 正确 Content-Type
    - 音视频 (.mp3/.wav/.mp4/.webm/.ogg)：StreamingResponse + Range 支持
    - Office (.docx/.xlsx/.pptx)：尝试 mammoth/openpyxl/python-pptx，缺失降级为下载链接
    - 文本类：返回 {type:text, content:...}
    - 其他：返回下载链接

    所有路径校验通过 _validate_file_path 防止路径遍历，并应用敏感文件 deny 列表。
    """
    # P0-4: RBAC 权限校验
    await _check_coding_permission(current_user, "coding:read", db)
    root = _resolve_coding_project_root(
        project_id=project_id,
        project_dir=project_dir,
        current_user=current_user,
        db=db,
        path_policy=path_policy,
    )
    normalized_project_id = project_id.strip() if project_id is not None else ""
    # 扩展名只取自客户端提供的安全相对路径，绝不暴露服务端项目根。
    ext = Path(path).suffix.lower()
    file_handle = _open_project_binary_file(path, root)

    # Markdown 渲染
    if ext in (".md", ".markdown"):
        return _render_markdown_file(file_handle)

    # 图片
    if ext in _IMAGE_MIME_MAP:
        return _stream_binary_file(
            file_handle, _IMAGE_MIME_MAP[ext], request, support_range=False
        )

    # 音频
    if ext in _AUDIO_MIME_MAP:
        return _stream_binary_file(
            file_handle, _AUDIO_MIME_MAP[ext], request, support_range=True
        )

    # 视频
    if ext in _VIDEO_MIME_MAP:
        return _stream_binary_file(
            file_handle, _VIDEO_MIME_MAP[ext], request, support_range=True
        )

    # Office 文件
    if ext in _OFFICE_EXTS:
        return _render_office_file(file_handle, ext, path, normalized_project_id)

    # 文本类
    if ext in _TEXT_MIME_MAP:
        try:
            content = file_handle.read().decode("utf-8")
            return {
                "type": "text",
                "content": content,
                "mime": _TEXT_MIME_MAP[ext],
            }
        except UnicodeDecodeError:
            # 编码失败，回退到下载
            return _download_link_response(path, normalized_project_id)
        finally:
            file_handle.close()

    # 未知类型不消费文件内容，返回下载链接前必须关闭已验证句柄。
    file_handle.close()
    return _download_link_response(path, normalized_project_id)


@router.get("/download")
async def download_file(
    path: str,
    project_id: Optional[str] = None,
    project_dir: Optional[str] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    path_policy: WorkbenchPathPolicy = Depends(get_coding_workbench_path_policy),
) -> StreamingResponse:
    """下载工作台项目内经过二次路径校验的常规文件。"""
    await _check_coding_permission(current_user, "coding:read", db)
    root = _resolve_coding_project_root(
        project_id=project_id,
        project_dir=project_dir,
        current_user=current_user,
        db=db,
        path_policy=path_policy,
    )
    file_handle = _open_project_binary_file(path, root)
    filename = _url_quote(Path(path).name, safe="")
    return StreamingResponse(
        _iter_open_file(file_handle),
        media_type="application/octet-stream",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )
